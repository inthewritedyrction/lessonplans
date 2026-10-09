"""Build Week9_Day4_Slides.pptx from the Day 2 deck's master/layout (same theme, header, background)."""
import copy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

SRC = Path(__file__).resolve().parent
TEMPLATE = SRC.parents[3] / "exemplars" / "week09" / "Week9_Day2_Slides.pptx"
OUT = SRC.parent / "Week9_Day4_Slides.pptx"
FONT = "Century Gothic"

INK, MUTED, WHITE = "1E2A32", "5A6B75", "FFFFFF"
YEL, YEL_D, YEL_L, YEL_INK = "FFC63C", "B98100", "FFF1C9", "3A2A00"
BLUE, BLUE_D, BLUE_L = "3B74E8", "2A56B5", "DCE7FB"
PUR, PUR_D, PUR_L = "8B6BE8", "6746C7", "E8DEFF"
TEAL, TEAL_D, TEAL_L = "00A39B", "007A73", "C8F1EB"
COR, COR_D, COR_L = "FF6B4E", "DF442A", "FFE1D6"

prs = Presentation(TEMPLATE)
# The template's notes master has no placeholders, so new notes slides would have no notes body.
# Borrow the slide-image + notes-body placeholders from an existing Day 2 notes slide.
NOTES_SP = [copy.deepcopy(sh._element) for sh in prs.slides[1].notes_slide.shapes
            if sh.is_placeholder and sh.placeholder_format.idx in (0, 1)]
# Drop the Day 2 slides; keep master, layouts, theme.
sld_ids = prs.slides._sldIdLst
for sld in list(sld_ids):
    prs.part.drop_rel(sld.rId)
    sld_ids.remove(sld)
LAYOUT = next(l for l in prs.slide_layouts if l.name == "CONTENT")
for sh in LAYOUT.shapes:
    if sh.has_text_frame and "WEEK 9" in sh.text_frame.text:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                r.text = r.text.replace("DAY 2", "DAY 4").replace("THE DIARY OF ANNE FRANK", "NIGHT BY ELIE WIESEL")


def rgb(h):
    return RGBColor.from_string(h)


def rect(s, x, y, w, h, fill=None, line=None, lw=2.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, dash=False, radius=0.08):
    r = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        r.adjustments[0] = radius
    if fill:
        r.fill.solid(); r.fill.fore_color.rgb = rgb(fill)
    else:
        r.fill.background()
    if line:
        r.line.color.rgb = rgb(line); r.line.width = Pt(lw)
        if dash:
            r.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        r.line.fill.background()
    r.shadow.inherit = False
    return r


def text(s, x, y, w, h, paras, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, spacing=1.0):
    """paras: str | list of paragraphs; each paragraph is str or list of (text, {opts}) runs."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        runs = [(para, {})] if isinstance(para, str) else para
        for t, o in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = rgb(o.get("color", color))
    return tb


def B(t, **o):
    return (t, {"bold": True, **o})


def N(t, **o):
    return (t, o)


def new_slide(num, bg=None):
    s = prs.slides.add_slide(LAYOUT)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    if bg:
        s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(bg)
    text(s, 12.45, 7.0, 0.5, 0.3, str(num), size=10, color="9AA6AD", align=PP_ALIGN.RIGHT)
    return s


def header(s, icon, icol, title, sub, chip=None, chipfill=None, dark=False):
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.5), Inches(0.62), Inches(0.62), Inches(0.62))
    o.fill.solid(); o.fill.fore_color.rgb = rgb(WHITE)
    o.line.color.rgb = rgb(icol); o.line.width = Pt(3)
    o.shadow.inherit = False
    text(s, 0.5, 0.62, 0.62, 0.62, icon, size=18, bold=True, color=icol, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.3, 0.6, 8.1, 0.5, title, size=28, bold=True, color=WHITE if dark else INK)
    text(s, 1.3, 1.1, 10.5, 0.35, sub, size=14, color="C9D4DA" if dark else MUTED)
    if chip:
        rect(s, 9.48, 0.62, 3.35, 0.38, fill=chipfill)
        text(s, 9.48, 0.62, 3.35, 0.38, chip, size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def pill(s, x, y, w, label, fill, color=WHITE, h=0.38, size=11):
    rect(s, x, y, w, h, fill=fill, radius=0.5)
    text(s, x, y, w, h, label, size=size, bold=True, color=color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def dark_bar(s, y, h, paras, size=20):
    rect(s, 0.5, y, 12.33, h, fill=INK)
    text(s, 0.85, y, 11.7, h, paras, size=size, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)


def notes(s, t):
    ns = s.notes_slide
    if ns.notes_placeholder is None:
        for el in NOTES_SP:
            ns.shapes._spTree.append(copy.deepcopy(el))
    ns.notes_text_frame.text = t


# ---------------------------------------------------------------- 1 · Title
s = new_slide(1, bg=INK)
for i, c in enumerate([TEAL, BLUE, COR, YEL, PUR]):
    d = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.7 + 0.5 * i), Inches(0.75), Inches(0.34), Inches(0.34))
    d.fill.solid(); d.fill.fore_color.rgb = rgb(c); d.line.fill.background()
text(s, 0.7, 1.35, 9.0, 0.35, "WEEK 9 · DAY 4 · ENCOUNTERING EVIL", size=13, bold=True, color="9FB3BD")
text(s, 0.7, 1.8, 12.0, 1.9, ["Entering Night:", "Moishe’s Warning and the Symbol of Night"], size=40, bold=True, color=WHITE)
text(s, 0.7, 3.75, 11.8, 0.4, "Night by Elie Wiesel · pp. 3–22 · first read done last night", size=15, color="C9D4DA")
rect(s, 0.7, 4.5, 5.8, 2.3, fill=PUR)
text(s, 1.0, 4.7, 5.2, 0.3, "ESSENTIAL QUESTION", size=11, bold=True, color=WHITE)
text(s, 1.0, 5.05, 5.2, 1.65, "What does the community’s reaction to the warning reveal about human nature, and what might Wiesel’s repeated word “night” come to symbolize?", size=15.5, color=WHITE)
rect(s, 6.8, 4.5, 5.8, 2.3, fill=YEL)
text(s, 7.1, 4.7, 5.2, 0.3, "LEARNING TARGET", size=11, bold=True, color=YEL_INK)
text(s, 7.1, 5.05, 5.2, 1.65, "I can analyze how the people of Sighet respond to warnings and how Wiesel begins to use the word “night” as a symbol, and use evidence from the text to explain what those patterns reveal.", size=14.5, color=YEL_INK)
notes(s, "Up as scholars walk in. Classwork packet on desks, p. 1. Homework sheet and printed reading (PDF pages 18-36) on top of the desk.")

# ---------------------------------------------------------------- 2 · Do Now
s = new_slide(2)
header(s, "01", YEL_D, "Do Now", "Silent. On your own. Use last night’s reading.", "PACKET P. 1 · BOX 01", YEL_D)
rect(s, 0.5, 1.75, 12.33, 2.6, fill=YEL_L)
text(s, 0.85, 1.75, 11.7, 2.6, [[N("Moishe returns to Sighet to warn the community about what he witnessed. ")], [B("What does the community do with his warning, and WHY do you think they respond that way?")], [N("Use one detail from last night’s reading.", size=19)]], size=24, anchor=MSO_ANCHOR.MIDDLE, spacing=1.08)
rect(s, 0.5, 4.65, 6.0, 1.4, fill=WHITE, line=YEL, lw=2.5)
text(s, 0.85, 4.65, 5.4, 1.4, [[B("Then circle:  ")], [N("Homework completed?  YES  /  NO")]], size=20, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 6.83, 4.65, 6.0, 1.4, fill=INK)
text(s, 7.15, 4.65, 5.4, 1.4, "Homework sheet + printed reading on top of your desk.", size=18, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
pill(s, 0.5, 6.35, 2.4, "5:00 · SILENT", YEL, YEL_INK)
notes(s, "SAY: Alright, look here. Homework sheet and your reading on top of your desk. Box 01. Moishe comes back to Sighet to warn them. What does the community do with his warning, and WHY do you think they respond that way? One detail from last night. Five minutes. Silent. Then circle YES or NO.\n"
         "CHECK while circulating: read? understood? evidence? inferring? Note 2-3 strong 'why' answers to call on.\n"
         "IF STALLED: Open to p. 7. Find what the people say about Moishe.\n"
         "TRANSITION: Pencils down. Homework sheet out. We're going to put last night's reading back together, fast. Fix or add to your answers while we talk.")

# ---------------------------------------------------------------- 3 · Homework debrief
s = new_slide(3)
header(s, "↺", BLUE_D, "Last Night’s Reading", "Fix or add to your homework while we talk.", "HOMEWORK · pp. 3–22", BLUE)
qs = ["What happened to Moishe?", "What did he come back trying to do?", "How did the people respond?", "WHY do you think they responded that way?"]
for i, q in enumerate(qs):
    y = 1.75 + i * 0.68
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.5), Inches(y + 0.06), Inches(0.46), Inches(0.46))
    o.fill.solid(); o.fill.fore_color.rgb = rgb(BLUE); o.line.fill.background()
    text(s, 0.5, y + 0.06, 0.46, 0.46, str(i + 1), size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.1, y, 11.5, 0.6, q, size=21, bold=(i == 3), anchor=MSO_ANCHOR.MIDDLE)
steps = [("Moishe’s warning", BLUE), ("Danger grows outside", BLUE), ("German soldiers arrive", PUR), ("Restrictions", PUR), ("Ghettos", COR)]
for i, (t, c) in enumerate(steps):
    x = 0.5 + i * 2.5
    rect(s, x, 4.75, 2.25, 1.0, fill=c)
    text(s, x, 4.75, 2.25, 1.0, t, size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if i < 4:
        text(s, x + 2.22, 4.75, 0.3, 1.0, "→", size=20, bold=True, color=INK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.5, 5.95, 12.33, 0.8, "Every step is closer. Watch what the town does at every step.", size=20, bold=True, color=MUTED)
notes(s, "Ask in order; 2 answers each with page numbers. 1) What happened to Moishe? (expelled; Gestapo shooting in the forest; wounded, left for dead, p. 6) 2) What did he come back trying to do? ('I wanted to come back to warn you,' p. 7) 3) How did the people respond? (refused to believe, refused to listen, pity, imagining things, 'gone mad,' p. 7) 4) WHY? Take 2-3, don't judge yet: 'Hold onto those.'\n"
         "IF 'they were dumb': That's a judgment. What does the text show about why?\n"
         "SEQUENCE (homework #5): Moishe's warning -> danger grows outside (Hungary) -> German soldiers arrive -> restrictions -> ghettos. Ten minutes max; don't reteach every detail.\n"
         "TRANSITION: Before we go back into the text, notice what kind of book this is. Yesterday we finished a play. This is something different. Box 02.")

# ---------------------------------------------------------------- 4 · Play vs Memoir
s = new_slide(4)
header(s, "02", BLUE_D, "From Play to Memoir", "Same history. A different way of telling it.", "PACKET P. 1 · BOX 02", BLUE)
rect(s, 0.5, 1.75, 6.0, 3.0, fill="EEF1F3")
text(s, 0.85, 1.9, 5.4, 0.5, "PLAY", size=24, bold=True, color=MUTED)
text(s, 0.85, 2.5, 5.4, 2.1, "Playwrights shape a story for an audience. We experience characters from the outside: dialogue, stage directions, performance.", size=18)
rect(s, 6.83, 1.75, 6.0, 3.0, fill=BLUE_L)
text(s, 7.18, 1.9, 5.4, 0.5, "MEMOIR", size=24, bold=True, color=BLUE_D)
text(s, 7.18, 2.5, 5.4, 2.1, "A survivor tells his own lived experience through memory and purpose. We are inside the narrator’s perspective.", size=18)
dark_bar(s, 5.05, 1.6, [[N("What can a memoir give us that a play "), B("cannot", color=YEL), N("?")]], size=26)
notes(s, "SAY (Prather): In a play, playwrights shape a story for an audience. We experienced Anne from the outside. A memoir is a survivor telling his own lived experience, through memory and purpose. We are inside his perspective.\n"
         "ASK (Prather): What can a memoir give us that a play cannot? Take 2. LISTEN FOR: his thoughts; it really happened; his voice; he is looking back.\n"
         "IF INCOMPLETE: point to p. 11, '(Poor Father! Of what then did you die?)' - the survivor's voice looking back.\n"
         "WIESEL (1 min): Born in Sighet, the town you just read about. A teenager when these events happened. He survived, and his writing became testimony: a record, so people remember. Stop there.")

# ---------------------------------------------------------------- 5 · Two lenses
s = new_slide(5, bg=INK)
header(s, "◎", YEL, "Today’s Two Lenses", "Track two things in the opening of Night.", dark=True)
rect(s, 0.5, 1.8, 6.0, 3.2, fill=TEAL)
text(s, 0.85, 1.95, 5.4, 0.4, "LENS 1 · ABOUT PEOPLE", size=14, bold=True, color=WHITE)
text(s, 0.85, 2.45, 5.4, 2.4, "How does this town respond every time danger gets closer?", size=24, bold=True, color=WHITE)
rect(s, 6.83, 1.8, 6.0, 3.2, fill=PUR)
text(s, 7.18, 1.95, 5.4, 0.4, "LENS 2 · ABOUT LANGUAGE", size=14, bold=True, color=WHITE)
text(s, 7.18, 2.45, 5.4, 2.4, "What does Wiesel begin doing with the word “night”?", size=24, bold=True, color=WHITE)
text(s, 0.5, 5.4, 12.33, 1.2, "One is about people. One is about language. Both tell us what is coming.", size=26, bold=True, color=YEL, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "SAY (Prather): Track two things: how this town responds every time danger gets closer, and what Wiesel begins doing with the word night. One is about people. One is about language. Both tell us what is coming.\n"
         "Don't pre-summarize the rest of the memoir.\n"
         "TRANSITION: Lens one first. Printed reading, p. 5. We're rereading Moishe's return slowly, because this is the moment the whole opening turns on.")

# ---------------------------------------------------------------- 6 · Moishe reread
s = new_slide(6)
header(s, "W", TEAL_D, "Reread · Moishe’s Warning", "Printed reading pp. 5–7. Read slowly. Mark the margin.", "READING pp. 5–7", TEAL)
marks = [("W", "a warning"), ("R", "a reaction"), ("?", "something to discuss")]
for i, (m, d) in enumerate(marks):
    x = 0.5 + i * 4.15
    rect(s, x, 1.75, 4.0, 1.5, fill=TEAL_L)
    text(s, x + 0.3, 1.75, 1.0, 1.5, m, size=48, bold=True, color=TEAL_D, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 1.3, 1.75, 2.6, 1.5, d, size=20, bold=True, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 0.5, 3.5, 12.33, 1.5, fill=WHITE, line=TEAL, lw=2)
text(s, 0.85, 3.5, 11.7, 1.5, "“I wanted to come back to warn you. Only no one is listening to me…”  (p. 7)", size=24, bold=True, anchor=MSO_ANCHOR.MIDDLE)
dark_bar(s, 5.3, 1.35, [[N("Moishe is "), B("the witness nobody believes.", color=YEL), N("  What changed in him? Where are your R’s?")]], size=20)
notes(s, "SAY: As I read, mark the margin. W for a warning. R for a reaction. ? for something you want to talk about. Three marks.\n"
         "READ ALOUD: from p. 5 ('And Moishe the Beadle, the poorest of the poor of Sighet...') through p. 6 (expulsion, return) to the bottom of p. 7 ('This was toward the end of 1942.'). Slow down on p. 7: 'Moishe was not the same. The joy in his eyes was gone.' Pause after 'Only no one is listening to me...'\n"
         "ASK: What changed in Moishe? Where did you put your R's?\n"
         "LISTEN FOR: he stopped singing and talking about God; only talks about what he saw. R's: 'refused to believe... refused to listen,' pity, imagining things, 'gone mad' (p. 7). Also p. 6: 'The deportees were quickly forgotten.'\n"
         "GROUNDING: His warning is end of 1942; the Germans reach Sighet in spring 1944. There was time.")

# ---------------------------------------------------------------- 7 · Warning Reactions model
s = new_slide(7)
header(s, "M", TEAL_D, "Warning Reactions · Model Row", "Moishe · pp. 6–7", "PACKET P. 2 · ROW M", TEAL)
cols = [("THE WARNING", "Moishe’s eyewitness account of the shootings (p. 6). He came back to warn them (p. 7).", TEAL_L, TEAL_D),
        ("THE REACTION", "They refused to believe and refused to listen. Pity. “Imagining things.” “Gone mad.” (p. 7)", TEAL_L, TEAL_D),
        ("THE INFERENCE", "What does this reveal about people?", WHITE, COR_D)]
for i, (h, d, f, c) in enumerate(cols):
    x = 0.5 + i * 4.15
    rect(s, x, 1.75, 4.0, 3.2, fill=f, line=(COR if i == 2 else None), lw=2.5, dash=(i == 2))
    text(s, x + 0.25, 1.9, 3.5, 0.4, h, size=14, bold=True, color=c)
    text(s, x + 0.25, 2.4, 3.5, 2.4, d, size=18, bold=(i == 2))
dark_bar(s, 5.25, 1.45, "The third column is not a summary. It answers: what does this reaction reveal about what people do with unbearable news?", size=19)
notes(s, "MODEL (doc cam on Row M): Column one, the warning: Moishe's eyewitness account, p. 6, and why he came back, p. 7. Column two, the reaction: refused to believe, refused to listen, called him mad, p. 7. Those two are already in your packet.\n"
         "SAY (Prather): The third column is not a summary. It answers: what does this reaction reveal about what people do with unbearable news?\n"
         "CONNECT: We already studied disbelief and indifference as conditions that helped persecution grow. Now Wiesel is showing us what that looks like inside one town.\n"
         "TRANSITION: Don't write column three yet. Purple strip, right under Row M. I want to hear how you think.")

# ---------------------------------------------------------------- 8 · LEAF
s = new_slide(8, bg=PUR_L)
header(s, "★", PUR_D, "Show Your Thinking", "Think first. Then say it.", "PACKET P. 2 · PURPLE STRIP", PUR)
rect(s, 0.5, 1.75, 12.33, 1.5, fill=WHITE, line=PUR, lw=2.5, dash=True)
text(s, 0.85, 1.75, 11.7, 1.5, [[N("Okay, now I need the inference. ")], [B("What does their reaction reveal about people?", size=28)]], size=20, anchor=MSO_ANCHOR.MIDDLE)
pill(s, 0.5, 3.5, 3.4, "THINK SILENTLY · 30–45 SEC", PUR)
for i, stem in enumerate(["First I noticed…", "Then I noticed the people…", "So I decided this reveals…", "Because…"]):
    x = 0.5 + i * 3.13
    rect(s, x, 4.1, 2.95, 1.1, fill=PUR)
    text(s, x + 0.1, 4.1, 2.75, 1.1, stem, size=17, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
dark_bar(s, 5.5, 1.2, [[B("BUILD · REVISE · CHALLENGE", color=YEL), N("     “I agree / would revise that because ______.”", bold=False)]], size=19)
notes(s, "LEAF ACTION STEP - right after scholars name how the town responded to Moishe, before anyone writes column 3.\n"
         "1) ASK: Okay, now I need the inference. What does their reaction reveal about people? 2) Silent 30-45 sec, thumb on chest.\n"
         "3) Scholar 1: First I noticed ___. Then I noticed the people ___. So I decided this reveals ___ because ___.\n"
         "PRESS, DON'T ANSWER: What detail made you think that? Are you explaining what happened, or what it reveals? What does this tell us about how people respond to unbearable information? Stay with the same scholar.\n"
         "4) Scholar 2: I agree / would revise that because ___.\n"
         "5) Draft one inference sentence together. LAND IT: Good. That difference matters today. We are not just tracking what happens. We are asking what the reactions reveal.")

# ---------------------------------------------------------------- 9 · Partner warnings
s = new_slide(9)
header(s, "03", TEAL_D, "Partner Warning Analysis", "Rows 2–3. You already read this. Select, don’t search.", "PACKET P. 2 · ROWS 2–3", TEAL)
rect(s, 0.5, 1.75, 7.7, 1.55, fill=TEAL_L)
text(s, 0.8, 1.85, 7.2, 1.4, [[B("ROW 2  ", color=TEAL_D), B("pp. 8–9")], [N("the war news, the Fascists take power, German troops enter Hungary, the Budapest report", size=15)]])
rect(s, 0.5, 3.5, 7.7, 1.55, fill=TEAL_L)
text(s, 0.8, 3.6, 7.2, 1.4, [[B("ROW 3  ", color=TEAL_D), B("pp. 9–12")], [N("German soldiers arrive, leaders arrested, new decrees, the ghettos", size=15)]])
rect(s, 0.5, 5.3, 7.7, 1.5, fill=WHITE, line=TEAL, lw=2)
text(s, 0.8, 5.35, 7.2, 1.4, [[B("Warning → reaction → ")], [N("what does it ", bold=False), B("reveal", color=COR_D), N(" about people?")]], size=19, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 8.45, 1.75, 4.38, 5.05, fill=YEL_L)
text(s, 8.75, 1.9, 4.0, 0.35, "THE BIG PRESSES", size=14, bold=True, color=YEL_D)
text(s, 8.75, 2.35, 3.9, 4.3, ["“You told me what they did. What does that reveal?”", "", "“What are they protecting themselves from having to believe?”", "", "“At what point does disbelief become dangerous?”"], size=16, spacing=1.05)
notes(s, "PARTNERS (34-44): Look on the flagged pages. Find the warning, find what the town does with it, then: what does it reveal? Talk it out. Each of you writes.\n"
         "Pace: 38 'Row 2 should have its inference.' 42 'You should be on Row 3's inference.' 44 'Finish your sentence.'\n"
         "Morehouse: Row 2 -> p. 9 'Yet we still were not worried.' Row 3 -> p. 10 the chocolates; p. 11 'It's not lethal.'\n"
         "GO OVER (44-50): For each row, one scholar gives the warning, one the reaction, one the inference, another builds/revises. Then the big presses. Do not answer them yourself.\n"
         "TRANSITION: That's lens one, people. Now lens two, language. Last night I only asked you to NOTICE a pattern. Today we figure out what the pattern might mean. Page 3.")

# ---------------------------------------------------------------- 10 · Symbolism + Night Tracker
s = new_slide(10)
header(s, "04", PUR_D, "Night Tracker · Find the Pattern", "Last night you noticed it. Today: what might it mean?", "PACKET P. 3", PUR)
rect(s, 0.5, 1.75, 12.33, 1.25, fill=PUR_L)
text(s, 0.85, 1.75, 11.7, 1.25, [[B("SYMBOLISM  ", color=PUR_D), N("a concrete image, word, object, or repeated detail that carries a larger idea beyond its literal meaning")]], size=19, anchor=MSO_ANCHOR.MIDDLE)
for i, (h, d) in enumerate([("1 · LITERAL", "What does the word mean right there?"), ("2 · WHAT’S HAPPENING", "What is going on in the scene around it?"), ("3 · PATTERN", "What keeps arriving with night?")]):
    x = 0.5 + i * 4.15
    rect(s, x, 3.25, 4.0, 1.6, fill=WHITE, line=PUR, lw=2)
    text(s, x + 0.25, 3.35, 3.5, 1.45, [[B(h, color=PUR_D, size=14)], [N(d, size=17)]], spacing=1.05)
rect(s, 0.5, 5.1, 12.33, 1.7, fill=INK)
text(s, 0.85, 5.1, 11.7, 1.7, [[N("Night keeps arriving with ", color=WHITE), B("______", color=YEL), N(", so Wiesel may be building “night” into a symbol of ", color=WHITE), B("______", color=YEL), N(".", color=WHITE)]], size=21, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "TEACH (50-56): Last night I only asked you to NOTICE the pattern. Today we figure out what the pattern might mean. Define symbolism.\n"
         "MODEL p. 12 'Night fell.': literally it got dark; around it, the normal evening breaks - father pulled into an emergency meeting, 'The story he had interrupted would remain unfinished.' Write: night comes when ordinary life gets interrupted, maybe the beginning of the end of safety. MAYBE - one row isn't a pattern.\n"
         "DO NOT post 'night = death/evil/fear.' Let the pattern emerge (Prather).\n"
         "ROWS (56-64): pp. 14, 18, 21, 22. Stuck: Don't jump to the symbol yet. What pattern do you actually see?\n"
         "GO OVER (64-69): What kinds of events keep happening when night appears? What pattern? What could night begin to represent beyond darkness? Accept multiple defensible readings. Then the pattern sentence.")

# ---------------------------------------------------------------- 11 · Focused write
s = new_slide(11)
header(s, "05", BLUE_D, "Focused Analytical Write", "Choose ONE. Talk first, then write.", "PACKET P. 4 · BOX 05", BLUE)
rect(s, 0.5, 1.75, 6.0, 2.2, fill=BLUE_L)
text(s, 0.8, 1.85, 5.5, 2.0, [[B("A · HUMAN NATURE", color=BLUE_D, size=14)], [N("What does the community’s response to the warnings reveal about human nature?", size=17)]], spacing=1.05)
rect(s, 6.83, 1.75, 6.0, 2.2, fill=PUR_L)
text(s, 7.13, 1.85, 5.5, 2.0, [[B("B · SYMBOLISM", color=PUR_D, size=14)], [N("What might Wiesel’s repeated use of “night” symbolize in the opening of the memoir?", size=17)]], spacing=1.05)
for i, (h, d) in enumerate([("CLAIM", "What are you arguing?"), ("EVIDENCE", "What line or detail proves it? Page?"), ("REASONING", "Why does that evidence support your claim?")]):
    x = 0.5 + i * 4.15
    rect(s, x, 4.2, 4.0, 1.25, fill=YEL_L)
    text(s, x + 0.25, 4.25, 3.5, 1.15, [[B(h, color=YEL_D, size=14)], [N(d, size=16)]])
dark_bar(s, 5.7, 1.1, "You found the line. Now tell me what it proves.", size=21)
notes(s, "SAY: Choose ONE. Prompt A if your warning chart is stronger, Prompt B if your night chart is. Claim, one strong piece of evidence with a page, reasoning. Done early? Add a second evidence and reasoning.\n"
         "TALK FIRST (1 min): say your claim to your partner in one sentence; partner says 'So what?' if it sounds like a summary.\n"
         "WHISPER ONLY: You found the line. What does it prove? / You gave me what happened. I need what it reveals.\n"
         "TIME: 8 minutes on the clock; 10-12 if the debriefs ran fast; finish at home if behind.\n"
         "TRANSITION: Pencils down. Turn to page 5. Exit ticket. On your own.")

# ---------------------------------------------------------------- 12 · Exit
s = new_slide(12)
header(s, "06", COR_D, "Exit Ticket · On Your Own", "Silent. Independent. Answer the one Ms. Briggs chooses.", "PACKET P. 5 · BOX 06", COR)
rect(s, 0.5, 1.75, 12.33, 2.0, fill=WHITE, line=COR, lw=2.5)
text(s, 0.85, 1.75, 11.7, 2.0, [[B("A  ", color=COR_D), N("One warning the people of Sighet fail to take seriously is ______. Their reaction reveals ______ about people.")]], size=21, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.5, 3.8, 12.33, 0.5, "OR", size=16, bold=True, color=COR_D, align=PP_ALIGN.CENTER)
rect(s, 0.5, 4.35, 12.33, 2.0, fill=WHITE, line=COR, lw=2.5)
text(s, 0.85, 4.35, 11.7, 2.0, [[B("B  ", color=COR_D), N("So far, Wiesel’s repeated use of “night” may symbolize ______ because ______.")]], size=21, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "CHOOSE: more time on warnings today -> A. More time on symbolism -> B. Say it once.\n"
         "Silent and independent. Collect exit tickets. Charts on pp. 2-3 stay in folders for Day 5's synthesis (Prather). Unfinished focused writes go home.\n"
         "IF BEHIND: the exit ticket still happens, even at two minutes.")

prs.save(OUT)
print("slides ok:", len(prs.slides))
