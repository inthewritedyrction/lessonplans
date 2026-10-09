"""Build Week9_Day3_Slides.pptx from the Day 2 deck's master/layout (same theme, header, background)."""
import copy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

SRC = Path(__file__).resolve().parent
TEMPLATE = SRC.parents[3] / "exemplars" / "week09" / "Week9_Day2_Slides.pptx"
OUT = SRC.parent / "Week9_Day3_Slides.pptx"
FONT = "Century Gothic"

INK, MUTED, WHITE = "1E2A32", "5A6B75", "FFFFFF"
YEL, YEL_D, YEL_L, YEL_INK = "FFC63C", "B98100", "FFF1C9", "3A2A00"
BLUE, BLUE_D, BLUE_L = "3B74E8", "2A56B5", "DCE7FB"
PUR, PUR_D, PUR_L = "8B6BE8", "6746C7", "E8DEFF"
TEAL, TEAL_D, TEAL_L = "00A39B", "007A73", "C8F1EB"
COR, COR_D, COR_L = "FF6B4E", "DF442A", "FFE1D6"

prs = Presentation(TEMPLATE)
# The template's notes master has no placeholders, so new notes slides would have no notes body.
# Borrow the slide-image + notes-body placeholders from an existing Day 2 notes slide.
NOTES_SP = [copy.deepcopy(sh._element) for sh in prs.slides[1].notes_slide.shapes
            if sh.is_placeholder and sh.placeholder_format.idx in (0, 1)]
# Drop the Day 2 slides; keep master, layouts, theme.
sld_ids = prs.slides._sldIdLst
for sld in list(sld_ids):
    prs.part.drop_rel(sld.rId)
    sld_ids.remove(sld)
LAYOUT = next(l for l in prs.slide_layouts if l.name == "CONTENT")
for sh in LAYOUT.shapes:
    if sh.has_text_frame and "WEEK 9" in sh.text_frame.text:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                r.text = r.text.replace("DAY 2", "DAY 3")


def rgb(h):
    return RGBColor.from_string(h)


def rect(s, x, y, w, h, fill=None, line=None, lw=2.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, dash=False, radius=0.08):
    r = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        r.adjustments[0] = radius
    if fill:
        r.fill.solid(); r.fill.fore_color.rgb = rgb(fill)
    else:
        r.fill.background()
    if line:
        r.line.color.rgb = rgb(line); r.line.width = Pt(lw)
        if dash:
            r.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        r.line.fill.background()
    r.shadow.inherit = False
    return r


def text(s, x, y, w, h, paras, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, spacing=1.0):
    """paras: str | list of paragraphs; each paragraph is str or list of (text, {opts}) runs."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        runs = [(para, {})] if isinstance(para, str) else para
        for t, o in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = rgb(o.get("color", color))
    return tb


def B(t, **o):
    return (t, {"bold": True, **o})


def N(t, **o):
    return (t, o)


def new_slide(num, bg=None):
    s = prs.slides.add_slide(LAYOUT)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    if bg:
        s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(bg)
    text(s, 12.45, 7.0, 0.5, 0.3, str(num), size=10, color="9AA6AD", align=PP_ALIGN.RIGHT)
    return s


def header(s, icon, icol, title, sub, chip=None, chipfill=None, dark=False):
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.5), Inches(0.62), Inches(0.62), Inches(0.62))
    o.fill.solid(); o.fill.fore_color.rgb = rgb(WHITE)
    o.line.color.rgb = rgb(icol); o.line.width = Pt(3)
    o.shadow.inherit = False
    text(s, 0.5, 0.62, 0.62, 0.62, icon, size=18, bold=True, color=icol, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.3, 0.6, 8.1, 0.5, title, size=28, bold=True, color=WHITE if dark else INK)
    text(s, 1.3, 1.1, 10.5, 0.35, sub, size=14, color="C9D4DA" if dark else MUTED)
    if chip:
        rect(s, 9.48, 0.62, 3.35, 0.38, fill=chipfill)
        text(s, 9.48, 0.62, 3.35, 0.38, chip, size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def pill(s, x, y, w, label, fill, color=WHITE, h=0.38, size=11):
    rect(s, x, y, w, h, fill=fill, radius=0.5)
    text(s, x, y, w, h, label, size=size, bold=True, color=color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def dark_bar(s, y, h, paras, size=20):
    rect(s, 0.5, y, 12.33, h, fill=INK)
    text(s, 0.85, y, 11.7, h, paras, size=size, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)


def notes(s, t):
    ns = s.notes_slide
    if ns.notes_placeholder is None:
        for el in NOTES_SP:
            ns.shapes._spTree.append(copy.deepcopy(el))
    ns.notes_text_frame.text = t


# ---------------------------------------------------------------- 1 · Title
s = new_slide(1, bg=INK)
for i, c in enumerate([TEAL, BLUE, COR, YEL, PUR]):
    d = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.7 + 0.5 * i), Inches(0.75), Inches(0.34), Inches(0.34))
    d.fill.solid(); d.fill.fore_color.rgb = rgb(c); d.line.fill.background()
text(s, 0.7, 1.35, 9.0, 0.35, "WEEK 9 · DAY 3 · ENCOUNTERING EVIL", size=13, bold=True, color="9FB3BD")
text(s, 0.7, 1.8, 12.0, 1.9, ["Theme and Legacy:", "Closing the Play"], size=44, bold=True, color=WHITE)
text(s, 0.7, 3.75, 11.8, 0.4, "The Diary of Anne Frank · Act II, Scenes 3–5 · pp. 766–778 · Course Mojo Lesson 6", size=15, color="C9D4DA")
rect(s, 0.7, 4.5, 5.8, 2.3, fill=PUR)
text(s, 1.0, 4.7, 5.2, 0.3, "ESSENTIAL QUESTION", size=11, bold=True, color=WHITE)
text(s, 1.0, 5.05, 5.2, 1.65, "How does the contrast between Anne’s tone and the mood of the final scene develop the theme of the play?", size=16.5, color=WHITE)
rect(s, 6.8, 4.5, 5.8, 2.3, fill=YEL)
text(s, 7.1, 4.7, 5.2, 0.3, "LEARNING TARGET", size=11, bold=True, color=YEL_INK)
text(s, 7.1, 5.05, 5.2, 1.65, "I can develop an arguable theme and explain how the contrast between Anne’s tone and the mood of the final scene supports that theme using specific evidence.", size=15, color=YEL_INK)
notes(s, "Up as scholars walk in. Packets on desks, page 1. Day 2 packets in folders: scholars need the Tone and Mood Map (Day 2, p. 2) at 5:00. Clip cued full screen.")

# ---------------------------------------------------------------- 2 · Do Now
s = new_slide(2)
header(s, "01", YEL_D, "Do Now · The Real Anne", "Watch twice: once to see, once to notice. Then write silently.", "PACKET P. 1 · BOX 01", YEL_D)
rect(s, 0.5, 1.75, 7.4, 2.75, fill=INK)
text(s, 0.85, 1.95, 6.8, 0.6, [[B("▶ ", color=YEL), B("“Anne Frank: the only existing film images”", color=WHITE)]], size=19)
text(s, 0.85, 2.6, 6.8, 0.7, "1941: the real Anne leans from a window to watch a wedding.", size=15, color="C9D4DA")
rect(s, 0.85, 3.45, 6.7, 0.75, line=COR, lw=2, dash=True)
text(s, 0.85, 3.45, 6.7, 0.75, "INSERT VIDEO LINK HERE · from the Prather Day 3 materials", size=13, bold=True, color=COR, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 8.15, 1.75, 4.68, 1.28, fill=YEL_L)
text(s, 8.45, 1.85, 4.2, 0.3, "WATCH 1", size=12, bold=True, color=YEL_D)
text(s, 8.45, 2.2, 4.2, 0.7, "Just see. Pencils down.", size=20, bold=True)
rect(s, 8.15, 3.22, 4.68, 1.28, fill=YEL_L)
text(s, 8.45, 3.32, 4.2, 0.3, "WATCH 2", size=12, bold=True, color=YEL_D)
text(s, 8.45, 3.67, 4.2, 0.7, "Notice. What do you catch now?", size=20, bold=True)
rect(s, 0.5, 4.8, 12.33, 1.3, fill=WHITE, line=YEL, lw=2.5)
text(s, 0.85, 4.8, 11.7, 1.3, [[N("You just finished the play. "), B("What does it change to see the real Anne Frank after finishing the play?")]], size=23, anchor=MSO_ANCHOR.MIDDLE)
pill(s, 0.5, 6.35, 3.6, "5:00 · WATCH TWICE, THEN SILENT", YEL, YEL_INK)
notes(s, "SAY: Alright, look here. Yesterday we finished the play. Before we talk about anything, I want you to see her. We're going to watch this twice. The first time, just see. Pencils down.\n"
         "Play once. SAY: Again. This time, notice. What do you see that you missed the first time? Play again.\n"
         "SAY: Pencils. Box 01. What does it change to see the real Anne Frank after finishing the play? Silent, on your own.\n"
         "Keep commentary minimal. Let the image work. This is not a history lecture. (Prather)\n"
         "IF STALLED: What did you notice the second time? Start there.\n"
         "TRANSITION (timer ends): Pencils down for a second. I want to hear a few of you. Just one line.")

# ---------------------------------------------------------------- 3 · Legacy
s = new_slide(3)
header(s, "◆", YEL_D, "Her Voice. Her Face.", "What does it change to see her? One line. Volunteers only.", "PACKET P. 1 · LEGACY", YEL_D)
rect(s, 0.5, 1.75, 12.33, 2.6, fill=INK)
text(s, 1.0, 1.75, 11.3, 2.6, ["“The play ends with her voice.", "This film is her face.", "Both are legacy: what remains and keeps speaking.”"], size=28, bold=True, color=YEL, anchor=MSO_ANCHOR.MIDDLE, spacing=1.1)
rect(s, 0.5, 4.65, 12.33, 1.5, fill=YEL_L)
text(s, 0.85, 4.75, 11.7, 0.35, "LEGACY", size=13, bold=True, color=YEL_D)
text(s, 0.85, 5.12, 11.7, 0.9, "what a person leaves behind that continues to matter after they are gone", size=24, bold=True)
text(s, 0.5, 6.35, 12.33, 0.4, "Hold onto that word. It’s on your exit ticket.", size=14, color=MUTED)
notes(s, "ASK: What does it change to see her? Volunteers only, one line each, 2-3 scholars. (Prather)\n"
         "LISTEN FOR: she's real, not just a character; a regular kid; the voice-over now has a face; the ending hurts more.\n"
         "IF: 'It's sad' -> Say more. What lands differently now than when we read the play?\n"
         "SAY (Prather): The play ends with her voice. This film is her face. Both are legacy: what remains and keeps speaking.\n"
         "Point to the legacy strip under Box 01.")

# ---------------------------------------------------------------- 4 · Retrieve Tuesday
s = new_slide(4)
header(s, "M", PUR_D, "Retrieve Tuesday · Tone vs. Mood", "Take out your Day 2 Tone and Mood Map. Row M.", "DAY 2 MAP · ROW M", PUR)
rect(s, 0.5, 1.75, 12.33, 1.0, fill=PUR_L)
text(s, 0.85, 1.75, 11.7, 1.0, "“In spite of everything, I still believe that people are really good at heart.”  (p. 778)", size=21, bold=True, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 0.5, 3.0, 6.0, 2.2, fill=BLUE_L)
text(s, 0.85, 3.15, 5.4, 0.4, "ANNE’S TONE = ______", size=22, bold=True, color=BLUE_D)
text(s, 0.85, 3.7, 5.4, 1.4, ["Her attitude. Proved by her diction.", [B("Which exact words?")]], size=17)
rect(s, 6.83, 3.0, 6.0, 2.2, fill=COR_L)
text(s, 7.18, 3.15, 5.4, 0.4, "AUDIENCE MOOD = ______", size=22, bold=True, color=COR_D)
text(s, 7.18, 3.7, 5.4, 1.4, ["Our feeling. Built by what we know.", [B("What does the audience know?")]], size=17)
dark_bar(s, 5.45, 1.35, [[N("Today’s question:  "), B("How does that contrast develop the play’s theme?", color=YEL)]], size=22)
notes(s, "SAY: Take out yesterday's Tone and Mood Map, page 2. Go to Row M.\n"
         "COLD CALL 2 (Prather): Restate the contrast in the final scene. Anne's tone is ___. The audience's mood is ___. What evidence did you use?\n"
         "LISTEN FOR: hopeful / defiant ('in spite of everything,' 'still believe,' 'really good at heart'). Mood: devastating; Mr. Frank: 'She'd been in Belsen with Anne... I know now.' (p. 778)\n"
         "IF: 'Anne's tone is sad' -> Whose feeling is 'sad,' Anne's or ours? Show me her actual words.\n"
         "POST under THEME on Tuesday's board: How does that contrast develop the play's theme? (Prather)\n"
         "TRANSITION: You can name the contrast. Today's question is what that contrast is for. To answer it, you need to know what a theme actually is. Box 02.")

# ---------------------------------------------------------------- 5 · Theme Ladder
s = new_slide(5)
header(s, "02", BLUE_D, "The Theme Ladder", "Climb from a topic to an arguable theme.", "PACKET PP. 1–2 · BOX 02", BLUE)
steps = [("1 · TOPIC", "one word", "hope", 4.35),
         ("2 · WHAT THE PLAY SAYS", "about this play", "The play is about holding onto hope.", 3.25),
         ("3 · ARGUABLE THEME", "about life or people beyond this play", "The playwrights argue that hope is both precious and unbearably fragile in the face of evil.", 1.75)]
for i, (lab, note, ex, y) in enumerate(steps):
    x = 0.5 + i * 0.55
    h = 0.95 if i < 2 else 1.4
    rect(s, x, y, 2.3, h, fill=BLUE)
    text(s, x + 0.15, y + 0.08, 2.05, h - 0.1, [[B(lab, size=14, color=WHITE)], [N(note, size=11, color=WHITE)]], size=14)
    rect(s, x + 2.4, y, 5.6 - i * 0.55 + 0.55 * 0, h, fill=BLUE_L)
    text(s, x + 2.6, y, 5.3 - i * 0.55, h, ex, size=17 if i < 2 else 17, italic=True, color=BLUE_D, bold=(i == 2), anchor=MSO_ANCHOR.MIDDLE)
rect(s, 9.35, 1.75, 3.48, 3.55, fill=WHITE, line=COR, lw=2)
text(s, 9.6, 1.85, 3.1, 0.35, "NOT A THEME", size=13, bold=True, color=COR_D)
text(s, 9.6, 2.25, 3.1, 3.0, [[B("Topic  ", color=COR_D), N("hope")], [B("Plot summary  ", color=COR_D), N("what happens in this play")], [B("Advice  ", color=COR_D), N("“Never give up hope.”")], [N(" ")], [B("“‘Never give up hope’ is a bumper sticker.”")]], size=15, spacing=1.15)
dark_bar(s, 5.55, 1.25, "A theme is an arguable idea about life or people beyond this one play, developed by the text and proved with evidence.", size=19)
notes(s, "SAY: Three things that are not theme. A topic is one word. A plot summary tells me what happened. Advice tells me what to do.\n"
         "MODEL THE CLIMB (Prather): Step one, topic: hope. Step two, what the play says: 'The play is about holding onto hope.' Step three: 'The playwrights argue that hope is both precious and unbearably fragile in the face of evil.' Somebody could disagree with that. That's how I know it's arguable.\n"
         "THEME VS. MORAL (Prather): 'Never give up hope' is a bumper sticker. That sounds nice, but nice is not enough. Can the text actually prove it? Compare: 'the playwrights suggest that hope survives even those it cannot save.'\n"
         "Scholars follow the model column. They don't fill My Ladder yet.")

# ---------------------------------------------------------------- 6 · Contrast -> Theme
s = new_slide(6, bg=INK)
header(s, "→", YEL, "Contrast → Theme", "Where does the theme live?", dark=True)
rect(s, 0.5, 1.8, 3.75, 1.75, fill=BLUE)
text(s, 0.75, 1.9, 3.3, 1.6, [[B("ANNE’S HOPEFUL TONE", size=13)], [N("alone → the theme would be simple", size=16)]], color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
text(s, 4.3, 1.8, 0.6, 1.75, "+", size=40, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 4.95, 1.8, 3.75, 1.75, fill=COR)
text(s, 5.2, 1.9, 3.3, 1.6, [[B("THE DEVASTATING MOOD", size=13)], [N("alone → the theme would be despair", size=16)]], color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
text(s, 8.75, 1.8, 0.6, 1.75, "→", size=40, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 9.4, 1.8, 3.43, 1.75, fill=YEL)
text(s, 9.6, 1.9, 3.0, 1.6, [[B("THE COLLISION", size=13)], [N("→ THEME", size=22, bold=True)]], color=YEL_INK, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 0.5, 3.9, 12.33, 1.45, fill="2A3A44")
text(s, 0.85, 3.9, 11.7, 1.45, "“The playwrights force the two together, and the theme lives in that collision.”", size=26, bold=True, color=YEL, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
text(s, 0.5, 5.65, 12.33, 0.9, "What idea about people survives the contrast?", size=30, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "SAY (Prather): If Anne's hopeful tone stood alone, the theme would be simple. If the devastating mood stood alone, the theme would be despair. The playwrights force the two together, and the theme lives in that collision.\n"
         "ASK (Prather): What idea about people survives the contrast? Wait. Take one or two. Do NOT answer it yourself. 'Hold that. Don't write it yet.'\n"
         "TRANSITION: That was my climb. Now I want to hear yours. Purple strip, right under the dark bar.")

# ---------------------------------------------------------------- 7 · LEAF
s = new_slide(7, bg=PUR_L)
header(s, "★", PUR_D, "Show Your Thinking", "Think first. Then say it.", "PACKET P. 2 · PURPLE STRIP", PUR)
rect(s, 0.5, 1.75, 12.33, 1.75, fill=WHITE, line=PUR, lw=2.5, dash=True)
text(s, 0.85, 1.85, 11.7, 1.6, [[N("Candidate:  "), B("“Never give up hope.”", size=30)], [N("Is it a topic, advice, or an arguable idea? How would you turn it into a real theme statement?", size=18)]], size=22, anchor=MSO_ANCHOR.MIDDLE, spacing=1.1)
pill(s, 0.5, 3.75, 3.4, "THINK SILENTLY · 30–45 SEC", PUR)
for i, stem in enumerate(["First I…", "Then I noticed…", "So I decided…", "Because…"]):
    x = 0.5 + i * 3.13
    rect(s, x, 4.35, 2.95, 0.95, fill=PUR)
    text(s, x, 4.35, 2.95, 0.95, stem, size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
dark_bar(s, 5.6, 1.15, [[B("BUILD · REVISE · CHALLENGE", color=YEL), N("     “I would revise it to ______ because ______.”", bold=False)]], size=19)
notes(s, "LEAF ACTION STEP. Place it here: after the ladder model, before scholars build their own.\n"
         "1) Read the candidate together. 2) SAY: Don't write. Don't talk. Is it a topic, advice, or an arguable idea? Thumb on your chest when you've got your 'First I...'. Wait the full 30-45 seconds.\n"
         "3) [Name], talk us through how you would turn that into a real theme statement. Start with 'First I...'\n"
         "PRESS, DON'T ANSWER: Is that a topic, advice, or an arguable idea? Does it apply beyond Anne? Could a reasonable person disagree? Can Tuesday's tone/mood evidence actually support it? Stay with the same scholar.\n"
         "4) Who can build on, revise, or challenge that theme statement? 'I would revise it to ___ because ___.'\n"
         "5) RELEASE: Good. That's the process I want running in your head: name what it is, push it past this one play, make sure somebody could disagree, then check it against your evidence. Your ladder. Go.")

# ---------------------------------------------------------------- 8 · Theme Test
s = new_slide(8)
header(s, "02", BLUE_D, "Build It. Test It.", "My Ladder → test it → backup candidate · 3 minutes", "PACKET P. 2 · THEME TEST", BLUE)
tests = ["A reasonable person could disagree.", "It applies beyond this one play.", "Tuesday’s tone/mood contrast can support it.", "It is an idea, not advice."]
for i, t in enumerate(tests):
    x, y = 0.5 + (i % 2) * 3.95, 1.75 + (i // 2) * 1.75
    rect(s, x, y, 3.75, 1.55, fill=BLUE_L)
    text(s, x + 0.25, y + 0.12, 3.3, 1.35, [[B("☐  ", size=22, color=BLUE_D)], [B(t)]], size=17)
rect(s, 8.45, 1.75, 4.38, 3.3, fill=YEL_L)
text(s, 8.75, 1.9, 3.9, 0.35, "OUR CANDIDATES", size=14, bold=True, color=YEL_D)
text(s, 8.75, 2.35, 3.9, 2.6, ["Charted on the board, word for word.", "Tested out loud.", [B("At least two stay up for Course Mojo.")]], size=16, spacing=1.15)
dark_bar(s, 5.35, 1.4, "That sounds nice, but can the text actually prove it?", size=22)
notes(s, "WRITE/WORK 3 min: My Ladder, check all four, write a backup candidate. Day 2 map open beside it.\n"
         "SAY: Go back to yesterday's map. What evidence do you already have?\n"
         "IF: topic -> Okay, that's the topic. I need the theme. Advice -> You gave me advice. Give me an idea about people. Plot -> That tells me what happens in this play. How can we turn that into an idea about people beyond Anne?\n"
         "CHART 2 min (Prather): take 2-3 scholar statements word for word. Test each: Is it arguable? Is it about life beyond this play? Can the tone/mood contrast serve as its evidence?\n"
         "Do NOT crown one correct theme. Leave at least two viable candidates posted for Mojo.\n"
         "TRANSITION: Those stay up. Before we analyze this ending in Course Mojo, I want you to sit with it for a few minutes. Page 3.")

# ---------------------------------------------------------------- 9 · Choice Board
s = new_slide(9)
header(s, "03", TEAL_D, "Closing the Play · Choice Board", "You have finished the play. Choose ONE task. 10 minutes.", "PACKET P. 3 · BOX 03", TEAL)
opts = [("1", "Found Poem", "6–10 exact lines arranged to capture the contrast"),
        ("2", "Tone Map", "hope rising and falling, turning points labeled with quote + page"),
        ("3", "Six-Word Memoirs", "Anne · Mr. Frank · the audience, each with its line"),
        ("4", "Letter Never Sent", "from Peter or Margot, grounded in one cited line"),
        ("5", "Director’s Note", "lighting, sound, audience focus for the final voice-over"),
        ("6", "Headline + Lede", "a 1947 review: headline + two sentences, one quoted line")]
for i, (n, h, d) in enumerate(opts):
    x, y = 0.5 + (i % 3) * 4.15, 1.75 + (i // 3) * 1.85
    rect(s, x, y, 4.0, 1.68, fill=TEAL_L)
    text(s, x + 0.2, y + 0.12, 3.65, 1.5, [[B(n + "  ", color=TEAL_D, size=20), B(h, size=18)], [N(d, size=14)]], spacing=1.05)
rect(s, 0.5, 5.6, 12.33, 1.15, fill=COR)
text(s, 0.85, 5.6, 11.7, 1.15, [[B("EVERY OPTION: at least one exact line + page number.  "), N("Start with your Day 2 map.", bold=False)]], size=21, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "SAY (Prather): You have finished The Diary of Anne Frank. Before we analyze its theme in Course Mojo, choose ONE task and complete it in the next ten minutes. Every option must include at least one exact line from the play with a page number.\n"
         "STEER quietly (Prather): stronger writers -> Option 1 or 5, two cited lines. Easier entry -> Option 2 or 3 with the starter stems. Same evidence requirement for everyone.\n"
         "PACE: 28 'Your line and page should be written.' 33 'Two minutes. Check your page number.' 35 'Pencils down.'\n"
         "TIME RULE: First thing to cut if behind. Unfinished = homework only if you direct it. Never take time from Mojo.\n"
         "TRANSITION: Hold onto that feeling. Now we're going to prove what this ending does. Course Mojo, Lesson 6. Page 4. Sit with your partner.")

# ---------------------------------------------------------------- 10 · Course Mojo
s = new_slide(10)
header(s, "04", PUR_D, "Course Mojo · Lesson 6", "Part A with your partner → Part B discussion → three reads → write", "PACKET P. 4 · BOX 04", PUR)
rect(s, 0.5, 1.75, 7.6, 5.0, fill=PUR_L)
text(s, 0.85, 1.9, 7.0, 0.35, "TARGET TASK", size=13, bold=True, color=PUR_D)
text(s, 0.85, 2.3, 7.0, 4.35, "At the very end of the play, we hear Anne say in a voice-over, “In spite of everything, I still believe that people are really good at heart” (p. 778). What is Anne’s tone in these lines? How does the contrast between her tone and the mood of the scene develop the theme of the play? Provide specific evidence from the text to support your answer.", size=19, spacing=1.08)
flow = [("PART A · PARTNERS", "Select the BEST evidence from your Day 2 map. Don’t search again.", TEAL_L, TEAL_D),
        ("PART B · DISCUSS", "Tone word · mood + what we know · an arguable theme", BLUE_L, BLUE_D),
        ("3 READS", "1 together · 2 what is it asking? · 3 what do I need?", YEL_L, YEL_D)]
for i, (h, d, f, c) in enumerate(flow):
    y = 1.75 + i * 1.7
    rect(s, 8.35, y, 4.48, 1.55, fill=f)
    text(s, 8.6, y + 0.1, 4.05, 1.4, [[B(h, color=c, size=13)], [N(d, size=15)]], spacing=1.05)
notes(s, "PART A (35-45): Read the Target Task once aloud so partners know where they're headed. SAY: Your job is not to find ten new quotes. Your job is to select the evidence that best answers this question. Yesterday's map is your evidence bank. Talk it out together, but each of you writes.\n"
         "Use the Course Mojo Lesson 6 Driving Questions and Criteria for Teachers on the platform (Prather). Not reproduced here.\n"
         "PRESS: You named tone and mood. What does the contrast DEVELOP? / Which theme candidate actually fits this evidence? / Which piece of evidence is strongest?\n"
         "WHEN STUCK: Talk me through what you know so far. What did you notice first? Which part of your evidence is doing the most work? What idea about people does that contrast point toward?\n"
         "PART B (45-52, Prather): whole-class discussion before writing: connotation-precise tone, mood with audience knowledge, theme as an arguable sentence. Require an arguable theme before drafting.\n"
         "3-READ: Read 1 together aloud. Read 2 on your own: what is it asking? Read 3 on your own: what evidence and skills do I need?")

# ---------------------------------------------------------------- 11 · Independent Writing
s = new_slide(11)
header(s, "04", PUR_D, "On Your Own · Target Task", "Silent. Independent. Focused literary analysis, not a five-paragraph essay.", "PACKET P. 5 · RESPONSE", PUR)
rect(s, 0.5, 1.75, 7.0, 3.55, fill=WHITE, line=PUR, lw=2)
text(s, 0.8, 1.85, 6.5, 0.35, "MY RESPONSE MUST HAVE", size=13, bold=True, color=PUR_D)
text(s, 0.8, 2.25, 6.5, 3.0, ["☐  A tone word I can defend + Anne’s diction", "☐  The mood + what the audience knows", "☐  An arguable theme", [N("☐  "), B("How"), N(" the contrast develops the theme")], [N("☐  At least "), B("two"), N(" cited pieces of evidence")]], size=17, spacing=1.2)
rect(s, 7.75, 1.75, 5.08, 3.55, fill=YEL_L)
text(s, 8.05, 1.85, 4.6, 0.35, "MILESTONES CHECK", size=13, bold=True, color=YEL_D)
text(s, 8.05, 2.25, 4.6, 3.0, [[B("Claim: "), N("Is my theme arguable?")], [B("Evidence: "), N("Did I cite precise lines + pages?")], [B("Reasoning: "), N("Did I explain how the contrast develops the theme?")]], size=17, spacing=1.2)
dark_bar(s, 5.55, 1.2, "You told me what the line says. Now tell me how it develops your theme.", size=20)
notes(s, "SAY: Thirteen minutes. Silent and on your own. Your claim is your theme. Two pieces of evidence with page numbers. After each one, explain how that line helps the contrast develop your theme.\n"
         "If your class submits on Course Mojo, scholars type there; p. 4 is the plan, p. 5 the backup.\n"
         "WHISPER ONLY: You told me what the line says. How does that line help develop your theme? / You named the contrast. Now tell me the so what. / Could you point to Tuesday's map and prove that theme?\n"
         "Morehouse: point to the 'Need a start?' frame on p. 4; the response stays independent. Done early: stretch line on p. 5.\n"
         "PACE: 57 tone and mood down. 61 explaining how the contrast develops your theme. 64 finish your sentence.\n"
         "TRANSITION: Pencils down. Find the dashed line on page 5. Exit ticket. On your own.")

# ---------------------------------------------------------------- 12 · Exit + Night
s = new_slide(12)
header(s, "05", COR_D, "Exit Ticket · On Your Own", "Silent. Independent. Then we close the play.", "PACKET P. 5 · BOX 05", COR)
rect(s, 0.5, 1.75, 7.4, 5.0, fill=WHITE, line=COR, lw=2.5)
prompts = ["The theme of The Diary of Anne Frank is ______. (an arguable sentence)", "The evidence that best develops it is ______ (p. ___).", "Anne’s legacy, to me, is ______."]
for i, p in enumerate(prompts):
    y = 2.0 + i * 1.5
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.8), Inches(y), Inches(0.48), Inches(0.48))
    o.fill.solid(); o.fill.fore_color.rgb = rgb(COR); o.line.fill.background()
    text(s, 0.8, y, 0.48, 0.48, str(i + 1), size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.45, y - 0.05, 6.2, 1.3, p, size=19)
rect(s, 8.15, 1.75, 4.68, 2.35, fill=COR_L)
text(s, 8.45, 1.85, 4.2, 0.35, "LEGACY SHARE", size=13, bold=True, color=COR_D)
text(s, 8.45, 2.25, 4.2, 1.8, ["Stand and read only:", [B("“Anne’s legacy, to me, is…”")], "No commentary. Let them sit."], size=15, spacing=1.1)
rect(s, 8.15, 4.3, 4.68, 2.45, fill=INK)
text(s, 8.45, 4.4, 4.2, 0.35, "TOMORROW · NIGHT", size=13, bold=True, color=YEL)
text(s, 8.45, 4.8, 4.2, 1.9, "A memoir by survivor Elie Wiesel, who was your age when the events begin. Its first pages contain a warning a whole town refuses to hear.", size=14, color=WHITE, spacing=1.05)
notes(s, "EXIT (65-71): Three sentences. Your theme has to be an arguable sentence, not a topic and not advice. Silent, on your own. Don't coach answers. Collect exit tickets.\n"
         "LEGACY SHARE (71-75, Prather): 4-5 volunteers stand and read ONLY 'Anne's legacy, to me, is...' No commentary between readers. Let the sentences accumulate. Protect the seriousness.\n"
         "PREVIEW (Prather): Tomorrow we open our second text. Night is not a play; it is a memoir, written by a survivor named Elie Wiesel who was your age when the events begin. The first pages contain a warning that a whole town refuses to hear.\n"
         "Preview only. Don't start Thursday's content. Day 2 map + Theme Builder go back in folders (Day 5 synthesis uses every chart).\n"
         "IF BEHIND: the exit ticket always happens; take fewer legacy readers, keep it serious.")

prs.save(OUT)
print("slides ok:", len(prs.slides))
